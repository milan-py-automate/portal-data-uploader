# ============================================================
# Portal Data Tool - FINAL COMBINED (Google Colab)
# IMAP → ZIP → XLSB→XLSX → TRANSFORM → Google Sheet (Sheet1, नीचे जोड़ो)
# ============================================================

import os
import re
import sys
import time
import email
import struct
import getpass
import zipfile
import warnings
import subprocess
import xml.etree.ElementTree as ET
from collections import Counter
from email.header import decode_header
from email.utils import parsedate_to_datetime
from datetime import datetime, timedelta, timezone


def _pip_install(*pkgs):
    subprocess.check_call([sys.executable, "-m", "pip", "install", "-q", *pkgs])


try:
    import imapclient
except ImportError:
    _pip_install("imapclient")
    import imapclient

import pandas as pd
import gspread
from openpyxl import Workbook

# pandas की तारीख़ वाली बेकार चेतावनियाँ बंद
warnings.filterwarnings("ignore", message=".*Parsing dates.*")
warnings.filterwarnings("ignore", message=".*Could not infer format.*")

# ---------- कॉन्फ़िग ----------
IMAP_SERVER = "shalbyhospitals.icewarpcloud.in"
IMAP_PORT   = 993
EMAIL       = "fa30.corp@shalby.in"
SUBJECT_KEYWORD = "Portal Data"

DOWNLOAD_DIR    = "downloads"
EXTRACT_DIR     = "extracted"
CONVERT_DIR     = "converted"
TRANSFORMED_DIR = "transformed"      # ट्रांसफ़ॉर्म हुई बैकअप Excel फ़ाइलें
CREDS_FILE      = "credentials.json"

IST = timezone(timedelta(hours=5, minutes=30))
ALL_TEXT = True                      # XLSB→XLSX में सब कुछ टेक्स्ट रखना (ट्रांसफ़ॉर्म में नंबर/तारीख़ ठीक होंगे)

SAVE_TRANSFORMED_COPY = True         # ट्रांसफ़ॉर्म हुई Excel की बैकअप कॉपी सेव करनी है या नहीं
WRITE_HEADER_IF_EMPTY = True         # शीट पूरी खाली हो तो पहली पंक्ति में हेडर लिखना
VALUE_INPUT   = "USER_ENTERED"       # (पुराने कोड जैसा ही) शीट में वैल्यू कैसे जाए
CHUNK_SIZE    = 5000                 # एक बार में इतनी पंक्तियाँ शीट में भेजना
SLEEP_BETWEEN_UNITS = 2              # यूनिट्स के बीच सेकंड (Google API लिमिट से बचाव)

# ---------- ट्रांसफ़ॉर्मेशन कॉन्फ़िग (कोड 2 से) ----------
REQUIRED_COLUMNS = [
    "Admission No.",
    "DOA",
    "DOD",
    "Invoice Date",
    "Submission To SRIT Date",
]
DATE_COLUMNS    = ["DOA", "DOD", "Invoice Date", "Submission To SRIT Date"]
NUMERIC_COLUMNS = ["Sr No", "Net Bill Amount", "Spon Amount"]

UNIT_SHEETS = {
    "Indore":   "1L53r-Adww_UqWwkF-FI7GJTDG1aF5An5seI4wr5bR70",
    "Jabalpur": "1OodEyJHBXb31dbQi5QheLgZVNhgyfNLkB02uhCPqriI",
    "Jaipur":   "1XrMEI8IMa4iLYg42Ttfp_rpqtuBwaXrGzRnSpOZxW9o",
    "Krishna":  "1_xBC3dp47YQCibDtfE4vcIk85bqWOB7oK25HsONXOSE",
    "Mohali":   "1tkhSlXQEvn9e97OZaPWJGHYDjDOXEhOdn58EvoxZWyU",
    "Naroda":   "167rgUg8wISP7qZNbntSzO8T_IyqkUpMXrabsXx09wm4",
    "SG":       "1PWpJOPF07BRfguhVPCwJPkB7DLitoUKlDFFy4-0-kOA",
    "Surat":    "1Betva_Q9ltdL-gCT5Bx2Y9IPsEbmt1GIEDeUe3ww5MQ",
    "Vapi":     "1uwkuW2iHnGJB42vjBwj_1o2YdyQa1Bh32e9qxMCs5cA",
    "Vijay":    "1_8Me5oruzSG4XSI9DLJ29Bv_Xrd-tqsdksYkovrDAiM",
}
UNIT_LOOKUP = {k.lower(): k for k in UNIT_SHEETS}


def match_unit(name):
    """फ़ाइल के नाम से यूनिट पहचानो (छोटे/बड़े अक्षर का फ़र्क़ नहीं)"""
    return UNIT_LOOKUP.get(str(name).strip().lower())


def get_secret(name, prompt):
    """पासवर्ड कोड में नहीं लिखना: पहले Colab Secrets, नहीं तो छुपा हुआ इनपुट"""
    try:
        from google.colab import userdata
        val = userdata.get(name)
        if val:
            return val
    except Exception:
        pass
    return getpass.getpass(prompt)


# ============================================================
# भाग A: IMAP
# ============================================================

def get_today():
    return datetime.now(IST).date()


def decode_mime_header(value):
    if not value:
        return ""
    parts = decode_header(value)
    out = ""
    for text, enc in parts:
        if isinstance(text, bytes):
            out += text.decode(enc or 'utf-8', errors='ignore')
        else:
            out += text
    return out


def mail_date_ist(msg, internaldate=None):
    if internaldate is not None:
        dt = internaldate
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone(IST).date()
    date_hdr = msg.get('Date') if msg else None
    if date_hdr:
        try:
            dt = parsedate_to_datetime(date_hdr)
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=timezone.utc)
            return dt.astimezone(IST).date()
        except Exception:
            pass
    return None


def fetch_todays_zips():
    os.makedirs(DOWNLOAD_DIR, exist_ok=True)
    today = get_today()
    print(f"\n📅 आज की तारीख़: {today.strftime('%d/%m/%Y')}")

    password = get_secret("IMAP_PASSWORD", "📧 ईमेल का पासवर्ड डालें: ")

    imap = imapclient.IMAPClient(IMAP_SERVER, port=IMAP_PORT, ssl=True)
    downloaded = []
    try:
        imap.login(EMAIL, password)
        imap.select_folder("INBOX")

        # कल से खोजो (सर्वर की टाइमज़ोन का फ़र्क़ न पड़े), फिर नीचे IST तारीख़ से छाँटो
        since_str = (today - timedelta(days=1)).strftime("%d-%b-%Y")
        uids = imap.search(['SINCE', since_str, 'SUBJECT', SUBJECT_KEYWORD])
        print(f"🔎 सर्वर से {len(uids)} मेल मिले")

        for uid in uids:
            resp = imap.fetch([uid], ['INTERNALDATE', 'RFC822'])[uid]
            internaldate = resp.get(b'INTERNALDATE')
            raw = resp.get(b'RFC822')
            msg = email.message_from_bytes(raw)
            subject = decode_mime_header(msg.get('Subject', ''))
            mdate = mail_date_ist(msg, internaldate)

            if mdate != today:
                print(f"⏭️  छोड़ा ({mdate}): {subject}")
                continue

            print(f"✅ मैच ({mdate}): {subject}")

            for part in msg.walk():
                fname = part.get_filename()
                if not fname:
                    continue
                fname = decode_mime_header(fname)
                if fname.lower().endswith('.zip'):
                    filepath = os.path.join(DOWNLOAD_DIR, fname)
                    if filepath in downloaded:      # एक ही नाम की दो ज़िप हों तो ओवरराइट न हो
                        filepath = os.path.join(DOWNLOAD_DIR, f"{uid}_{fname}")
                    with open(filepath, 'wb') as f:
                        f.write(part.get_payload(decode=True))
                    print(f"   💾 डाउनलोड: {os.path.basename(filepath)}")
                    downloaded.append(filepath)
    finally:
        try:
            imap.logout()
        except Exception:
            pass

    print(f"\n📦 कुल {len(downloaded)} ज़िप डाउनलोड हुईं")
    return downloaded


def extract_zips(zip_files):
    os.makedirs(EXTRACT_DIR, exist_ok=True)
    all_files = []
    for zpath in zip_files:
        zip_name = os.path.basename(zpath)
        out_dir = os.path.join(EXTRACT_DIR, zip_name.replace('.zip', ''))
        os.makedirs(out_dir, exist_ok=True)
        with zipfile.ZipFile(zpath, 'r') as z:
            print(f"\n📦 {zip_name} → {len(z.namelist())} एंट्री")
            for name in z.namelist():
                if name.endswith('/') or '__MACOSX' in name:
                    continue
                z.extract(name, out_dir)
                full = os.path.join(out_dir, name)
                all_files.append(full)
                print(f"   📄 {name}")
    print(f"\n📊 कुल {len(all_files)} फ़ाइलें")
    return all_files


# ============================================================
# भाग B: .xlsb → .xlsx
# ============================================================

def records(data):
    pos, n = 0, len(data)
    while pos < n:
        b = data[pos]; pos += 1
        rid = b & 0x7F
        if b & 0x80:
            b2 = data[pos]; pos += 1
            rid |= (b2 & 0x7F) << 7
        size, shift = 0, 0
        for _ in range(4):
            b = data[pos]; pos += 1
            size |= (b & 0x7F) << shift
            shift += 7
            if not (b & 0x80):
                break
        yield rid, data[pos:pos + size]
        pos += size


def wide_str(buf, off):
    ln = struct.unpack_from("<I", buf, off)[0]
    s = buf[off + 4: off + 4 + 2 * ln].decode("utf-16-le", "replace")
    return s, off + 4 + 2 * ln


def read_sst(zf):
    out = []
    if "xl/sharedStrings.bin" not in zf.namelist():
        return out
    for rid, d in records(zf.read("xl/sharedStrings.bin")):
        if rid == 19:
            out.append(wide_str(d, 1)[0])
    return out


def read_sheets(zf):
    rels = {}
    root = ET.fromstring(zf.read("xl/_rels/workbook.bin.rels"))
    for el in root:
        rels[el.attrib["Id"]] = el.attrib["Target"]
    sheets = []
    for rid, d in records(zf.read("xl/workbook.bin")):
        if rid == 156:
            r, off = wide_str(d, 8)
            name, _ = wide_str(d, off)
            t = rels[r].lstrip("/")
            if not t.startswith("xl/"):
                t = "xl/" + t
            sheets.append((name, t))
    return sheets


def rk(v):
    if v & 2:
        x = struct.unpack("<i", struct.pack("<I", v & 0xFFFFFFFC))[0] >> 2
    else:
        x = struct.unpack("<d", struct.pack("<Q", (v & 0xFFFFFFFC) << 32))[0]
    if v & 1:
        x /= 100
    return x


def to_text(v):
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        if v == int(v) and abs(v) < 1e15:
            return str(int(v))
        return repr(v)
    if isinstance(v, int):
        return str(v)
    return v


def parse_sheet(data, sst, ws, stats):
    row, next_col = 0, 0
    for rid, d in records(data):
        if rid == 0:
            row = struct.unpack_from("<I", d, 0)[0]
            next_col = 0
            if len(d) >= 21 and struct.unpack_from("<I", d, 13)[0] > 0:
                next_col = struct.unpack_from("<I", d, 17)[0]
        elif rid == 176 and len(d) >= 16:
            r1, r2, c1, c2 = struct.unpack_from("<iiii", d, 0)
            ws.merge_cells(start_row=r1 + 1, start_column=c1 + 1,
                           end_row=r2 + 1, end_column=c2 + 1)
        elif 1 <= rid <= 18:
            if rid <= 11:
                col = struct.unpack_from("<I", d, 0)[0]
                p = 8
            else:
                col = next_col
                p = 4
            next_col = col + 1
            v = None
            try:
                if rid in (2, 13):   v = rk(struct.unpack_from("<I", d, p)[0])
                elif rid in (4, 15, 10): v = bool(d[p])
                elif rid in (5, 16, 9):  v = struct.unpack_from("<d", d, p)[0]
                elif rid in (6, 17, 8):  v = wide_str(d, p)[0]
                elif rid in (7, 18):     v = sst[struct.unpack_from("<I", d, p)[0]]
                elif rid in (3, 14, 11): v = "#ERR"
            except Exception:
                stats["bad"] += 1
            stats[rid] += 1
            if v is not None:
                if ALL_TEXT:
                    v = to_text(v)
                c = ws.cell(row=row + 1, column=col + 1, value=v)
                if ALL_TEXT:
                    c.number_format = "@"


def convert_xlsb_to_xlsx(src, dst):
    out = Workbook()
    out.remove(out.active)
    with zipfile.ZipFile(src) as zf:
        sst = read_sst(zf)
        for name, path in read_sheets(zf):
            ws = out.create_sheet(title=re.sub(r"[\[\]:*?/\\]", "_", name)[:31])
            stats = Counter()
            parse_sheet(zf.read(path), sst, ws, stats)
    out.save(dst)


def convert_all_xlsb(files):
    os.makedirs(CONVERT_DIR, exist_ok=True)
    converted = []
    for f in files:
        if not f.lower().endswith('.xlsb'):
            continue
        base = os.path.splitext(os.path.basename(f))[0]
        parent = os.path.basename(os.path.dirname(f))
        out_dir = os.path.join(CONVERT_DIR, parent)
        os.makedirs(out_dir, exist_ok=True)
        dst = os.path.join(out_dir, base + ".xlsx")
        try:
            convert_xlsb_to_xlsx(f, dst)
            converted.append((base, dst))
            print(f"   ✅ {base}.xlsx")
        except Exception as e:
            print(f"   ❌ {base}: {e}")
    print(f"\n🔄 कुल {len(converted)} फ़ाइलें तैयार")
    return converted


# ============================================================
# भाग C: डेटा ट्रांसफ़ॉर्मेशन (कोड 2 की सारी लॉजिक)
# ============================================================

def convert_date(value):
    """तारीख़ को असली तारीख़ में बदलो (टेक्स्ट, Excel सीरियल नंबर, या पहले से तारीख़)"""
    if pd.isna(value) or str(value).strip() == "":
        return pd.NaT

    # पहले से वास्तविक तारीख़ हो
    if isinstance(value, (datetime, pd.Timestamp)):
        return pd.to_datetime(value, errors="coerce")

    # Excel serial date (असली नंबर)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if 1 <= value <= 100000:
            return pd.to_datetime(value, unit="D", origin="1899-12-30", errors="coerce")
        return pd.NaT

    text = str(value).strip()

    # Excel serial date (टेक्स्ट के रूप में, जैसे "46310") - XLSB→XLSX के ALL_TEXT की वजह से
    try:
        num = float(text.replace(",", ""))
        if 1 <= num <= 100000:
            return pd.to_datetime(num, unit="D", origin="1899-12-30", errors="coerce")
    except ValueError:
        pass

    formats = [
        "%d-%b-%y",    # 15-OCT-26
        "%d-%b-%Y",
        "%d-%m-%Y",
        "%d-%m-%y",
        "%d/%m/%Y",
        "%d/%m/%y",
        "%Y-%m-%d",
        "%d-%B-%Y",
    ]
    for fmt in formats:
        try:
            return pd.Timestamp(datetime.strptime(text, fmt))
        except ValueError:
            continue

    return pd.to_datetime(text, errors="coerce", dayfirst=True)


def convert_to_number(value):
    """टेक्स्ट वाले नंबर को असली नंबर में बदलो"""
    if pd.isna(value):
        return value

    if isinstance(value, (int, float)):
        return value

    text = str(value).strip()
    if text == "":
        return None

    text = (text.replace(",", "")
                .replace("₹", "")
                .replace("Rs.", "")
                .replace("Rs", "")
                .strip())

    # (1200) को -1200 मानो
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1].strip()

    try:
        number = float(text)
        if number.is_integer():
            return int(number)
        return number
    except (ValueError, OverflowError):
        return value


def transform_dataframe(df):
    """एक यूनिट के डेटा पर पूरी ट्रांसफ़ॉर्मेशन; गड़बड़ हो तो ValueError"""
    df = df.copy()

    # कॉलम नाम साफ़ करो
    df.columns = [str(c).strip() if c is not None else "" for c in df.columns]

    # 1. ज़रूरी कॉलम जाँचो
    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        raise ValueError(f"ज़रूरी कॉलम नहीं मिले: {missing} | मौजूद: {df.columns.tolist()}")

    # 2. 'Submission To SRIT Date' के बाद के कॉलम हटाओ
    srit_index = df.columns.get_loc("Submission To SRIT Date")
    df = df.iloc[:, :srit_index + 1].copy()

    # 3. Claim ID का डेटा खाली करो
    if "Claim ID" in df.columns:
        df["Claim ID"] = ""
    else:
        print("      ⚠️  चेतावनी: Claim ID कॉलम नहीं मिला")

    # 4. File Type: OP / IP (पहले से भरी वैल्यू नहीं बदलनी)
    if "File Type" in df.columns:
        existing_type = df["File Type"].copy()
        df = df.drop(columns=["File Type"])
    else:
        existing_type = pd.Series([None] * len(df), index=df.index, dtype=object)

    admission = df["Admission No."].fillna("").astype(str).str.strip()
    calculated_type = admission.apply(lambda x: "OP" if "-" in x else "IP")

    existing_clean = existing_type.fillna("").astype(str).str.strip()
    is_blank = existing_clean.eq("")

    df["File Type"] = existing_type
    df.loc[is_blank, "File Type"] = calculated_type.loc[is_blank]

    # File Type को D कॉलम (index 3) में रखो
    cols = df.columns.tolist()
    cols.remove("File Type")
    cols.insert(3, "File Type")
    df = df[cols]

    # 5. 'Month -Year' कॉलम Invoice Date के ठीक बाद
    if "Month -Year" in df.columns:
        df = df.drop(columns=["Month -Year"])
    invoice_index = df.columns.get_loc("Invoice Date")
    df.insert(invoice_index + 1, "Month -Year", "")

    # 6. चारों तारीख़ कॉलम बदलो
    for col in DATE_COLUMNS:
        df[col] = df[col].apply(convert_date)

    # 7. Month -Year बनाओ (जैसे Oct-2026); Invoice Date खाली तो खाली
    df["Month -Year"] = df["Invoice Date"].apply(
        lambda d: d.strftime("%b-%Y") if pd.notna(d) else ""
    )

    # 8. टेक्स्ट नंबर → असली नंबर
    for col in NUMERIC_COLUMNS:
        if col in df.columns:
            df[col] = df[col].apply(convert_to_number)
        else:
            print(f"      ⚠️  ध्यान दें: '{col}' कॉलम नहीं मिला")

    # 9. Invoice Date से सॉर्ट: पुरानी → नई
    df = df.sort_values(
        by="Invoice Date", ascending=True, na_position="last", kind="stable"
    ).reset_index(drop=True)

    return df


def transform_all(converted):
    """हर यूनिट की xlsx पढ़कर ट्रांसफ़ॉर्म करो. एक में गड़बड़ हो तो बाक़ी चलते रहें"""
    print("\n" + "=" * 60)
    print("🛠️  डेटा ट्रांसफ़ॉर्म कर रहे हैं")
    print("=" * 60)

    if SAVE_TRANSFORMED_COPY:
        os.makedirs(TRANSFORMED_DIR, exist_ok=True)

    transformed, results = [], []
    for unit, xlsx_path in converted:
        _set_unit(unit)
        print(f"\n🧩 {unit}")
        try:
            df = pd.read_excel(
                xlsx_path, engine="openpyxl", sheet_name=0, dtype=object,
                keep_default_na=False, na_values=[""],   # "NA" जैसे असली टेक्स्ट गायब न हों
            )
            df = df.dropna(how="all").reset_index(drop=True)
            if df.empty:
                raise ValueError("फ़ाइल में कोई डेटा नहीं")

            tdf = transform_dataframe(df)
            print(f"      ✅ {len(tdf)} पंक्तियाँ, {len(tdf.columns)} कॉलम ट्रांसफ़ॉर्म हुए "
                  f"(OP: {(tdf['File Type'].astype(str).str.strip() == 'OP').sum()}, "
                  f"IP: {(tdf['File Type'].astype(str).str.strip() == 'IP').sum()}, "
                  f"खाली Invoice Date: {tdf['Invoice Date'].isna().sum()})")

            if SAVE_TRANSFORMED_COPY:
                out_path = os.path.join(TRANSFORMED_DIR, f"{unit}_Converted.xlsx")
                with pd.ExcelWriter(out_path, engine="openpyxl",
                                    datetime_format="DD-MM-YYYY") as writer:
                    tdf.to_excel(writer, index=False, sheet_name="Processed Data")

            transformed.append((unit, tdf))
        except Exception as e:
            print(f"      ❌ ट्रांसफ़ॉर्म फ़ेल: {e}")
            results.append((unit, "❌ ट्रांसफ़ॉर्म फ़ेल", str(e)))

    return transformed, results


# ============================================================
# भाग D: Google Sheet
# ============================================================

def google_client():
    if not os.path.exists(CREDS_FILE):
        try:
            from google.colab import files
            print(f"📁 {CREDS_FILE} अपलोड करें:")
            up = files.upload()
            if up:
                src = next(iter(up.keys()))
                if src != CREDS_FILE:
                    os.replace(src, CREDS_FILE)
        except ImportError:
            pass
    if not os.path.exists(CREDS_FILE):
        raise FileNotFoundError(f"{CREDS_FILE} नहीं मिली")
    return gspread.service_account(filename=CREDS_FILE)


def clean_value(v):
    """शीट में भेजने लायक सादी Python वैल्यू"""
    if hasattr(v, "item"):
        try:
            v = v.item()
        except Exception:
            pass
    if isinstance(v, float) and v != v:   # NaN
        return ""
    return v


def df_to_sheet_rows(df):
    """ट्रांसफ़ॉर्म हुए डेटा को शीट की पंक्तियों में बदलो"""
    out = df.copy()
    epoch = pd.Timestamp("1899-12-30")

    # तारीख़ें सीरियल नंबर बनकर जाएँगी (हर लोकेल में सही), बाद में dd-mm-yyyy फ़ॉर्मेट लगेगा
    for col in DATE_COLUMNS:
        out[col] = out[col].apply(
            lambda d: int((d.normalize() - epoch).days) if pd.notna(d) else ""
        )

    # 'Oct-2026' को शीट तारीख़ न बना दे, इसलिए शुरू में ' लगाया (टेक्स्ट बना रहेगा)
    out["Month -Year"] = out["Month -Year"].apply(lambda v: f"'{v}" if v else "")

    out = out.astype(object).where(out.notna(), "")
    return [[clean_value(v) for v in row] for row in out.values.tolist()]


def append_to_sheet(client, sheet_id, header, rows, date_col_idx):
    """हर बार जोड़ने से पहले एक खाली पंक्ति (सिपरेटर)"""
    if not rows:
        print("      ⚠️  कोई डेटा नहीं")
        return 0

    sh = client.open_by_key(sheet_id)
    ws = sh.sheet1
    existing = ws.get_all_values()

    if not existing:
        start_row = 2                      # पहली पंक्ति खाली (या हेडर)
        if WRITE_HEADER_IF_EMPTY:
            ws.update(values=[header], range_name="A1", value_input_option=VALUE_INPUT)
            print("      🏷️  शीट खाली थी, पहली पंक्ति में हेडर लिखा")
    else:
        start_row = len(existing) + 2      # सिपरेटर + डेटा

    end_row = start_row + len(rows) - 1

    # ग्रिड छोटी पड़े तो बढ़ाओ
    if end_row > ws.row_count:
        ws.add_rows(end_row - ws.row_count + 50)
    if len(header) > ws.col_count:
        ws.add_cols(len(header) - ws.col_count)

    for i in range(0, len(rows), CHUNK_SIZE):
        chunk = rows[i:i + CHUNK_SIZE]
        ws.update(values=chunk, range_name=f"A{start_row + i}",
                  value_input_option=VALUE_INPUT)

    # चारों तारीख़ कॉलम पर dd-mm-yyyy फ़ॉर्मेट (एक ही API कॉल में)
    if date_col_idx:
        requests = [{
            "repeatCell": {
                "range": {
                    "sheetId": ws.id,
                    "startRowIndex": start_row - 1,
                    "endRowIndex": end_row,
                    "startColumnIndex": c,
                    "endColumnIndex": c + 1,
                },
                "cell": {"userEnteredFormat": {
                    "numberFormat": {"type": "DATE", "pattern": "dd-mm-yyyy"}}},
                "fields": "userEnteredFormat.numberFormat",
            }
        } for c in date_col_idx]
        sh.batch_update({"requests": requests})

    print(f"      ✅ {len(rows)} पंक्तियाँ जोड़ीं "
          f"(खाली पंक्ति {start_row - 1} के बाद, डेटा {start_row}–{end_row})")
    return len(rows)


def push_to_sheets(transformed, results):
    print("\n" + "=" * 60)
    print("📤 ट्रांसफ़ॉर्म हुआ डेटा Google Sheet में डाल रहे हैं")
    print("=" * 60)

    client = google_client()

    for unit_raw, tdf in transformed:
        unit = match_unit(unit_raw)
        if unit is None:
            _set_unit(unit_raw)
            print(f"\n⏭️  {unit_raw}: कोई शीट नहीं मिली, छोड़ा")
            results.append((unit_raw, "⏭️ शीट नहीं मिली", ""))
            continue

        _set_unit(unit)
        print(f"\n📄 {unit} → शीट में जोड़ रहे हैं")
        try:
            header = [str(c) for c in tdf.columns]
            date_idx = [tdf.columns.get_loc(c) for c in DATE_COLUMNS]
            rows = df_to_sheet_rows(tdf)
            n = append_to_sheet(client, UNIT_SHEETS[unit], header, rows, date_idx)
            results.append((unit, "✅ पूरा", f"{n} पंक्तियाँ"))
        except Exception as e:
            print(f"      ❌ गड़बड़: {e}")
            results.append((unit, "❌ शीट में डालना फ़ेल", str(e)))

        time.sleep(SLEEP_BETWEEN_UNITS)


# ============================================================
# भाग F: डैशबोर्ड लॉगिंग (Vercel डैशबोर्ड के लिए)
# ============================================================
import io
import threading

LOG_SHEET_ID = "1htKgsJmBBBETCy4yUdOBOu9CETJLy37pxHuEPRPPecI"            # ← लॉग वाली Google Sheet की ID यहाँ डालें (service account को Editor दें)
LOG_TAB      = "logs"
FLUSH_EVERY  = 3             # सेकंड: बफ़र शीट में भेजना
BEAT_EVERY   = 60            # सेकंड: "ज़िंदा हूँ" हार्टबीट (अटकने की पहचान के लिए)
LG = None


class RunLogger:
    def __init__(self):
        self.run_id = datetime.now(IST).strftime("%Y%m%d-%H%M%S")
        self.stage, self.unit = "", ""
        self.buf, self.lock, self.ok = [], threading.Lock(), False
        self.stop = threading.Event()
        try:
            sh = google_client().open_by_key(LOG_SHEET_ID)
            try:
                self.ws = sh.worksheet(LOG_TAB)
            except gspread.WorksheetNotFound:
                self.ws = sh.add_worksheet(LOG_TAB, rows=2000, cols=6)
                self.ws.update(values=[["run_id", "ts", "stage", "unit", "level", "message"]],
                               range_name="A1")
            self.ok = True
            threading.Thread(target=self._loop, daemon=True).start()
        except Exception as e:
            print(f"⚠️  डैशबोर्ड लॉगिंग बंद ({e})", file=sys.__stdout__)

    def log(self, msg, level="INFO", stage=None, unit=None):
        if not self.ok:
            return
        row = [self.run_id, datetime.now(IST).isoformat(timespec="seconds"),
               self.stage if stage is None else stage,
               self.unit if unit is None else unit, level, str(msg)[:500]]
        with self.lock:
            self.buf.append(row)

    def flush(self):
        with self.lock:
            rows, self.buf = self.buf, []
        if rows:
            try:
                self.ws.append_rows(rows, value_input_option="RAW")
            except Exception:
                with self.lock:
                    self.buf = (rows + self.buf)[-2000:]

    def _loop(self):
        last_beat = 0
        while not self.stop.wait(FLUSH_EVERY):
            if time.time() - last_beat >= BEAT_EVERY:
                self.log("heartbeat", "BEAT")
                last_beat = time.time()
            self.flush()

    def close(self):
        self.stop.set()
        self.flush()


class Tee(io.TextIOBase):
    """हर print() की लाइन स्क्रीन पर भी और डैशबोर्ड लॉग में भी"""
    def __init__(self, orig, lg):
        self.orig, self.lg, self.partial = orig, lg, ""

    def write(self, s):
        self.orig.write(s)
        self.partial += s
        *lines, self.partial = self.partial.split("\n")
        for ln in lines:
            t = ln.strip()
            if not t or set(t) <= {"="}:
                continue
            lvl = ("ERROR" if "❌" in t else
                   "WARN" if ("⚠️" in t or "⏭️" in t) else
                   "SUCCESS" if "✅" in t else "INFO")
            self.lg.log(t, lvl)
        return len(s)

    def flush(self):
        self.orig.flush()

    def isatty(self):
        return False


def _stage(name):
    if LG:
        LG.stage, LG.unit = name, ""


def _set_unit(name):
    if LG:
        LG.unit = name


# ============================================================
# भाग E: मुख्य
# ============================================================

def _run_inner():
    print("=" * 60)
    print("🚀 Portal Data Tool (Colab) - Combined")
    print("=" * 60)

    _stage("IMAP")
    zips = fetch_todays_zips()
    if not zips:
        print("\n⚠️  कोई ज़िप नहीं मिली।")
        return

    _stage("ZIP")
    files_ = extract_zips(zips)
    if not files_:
        print("\n⚠️  ज़िप में कुछ नहीं मिला।")
        return

    _stage("CONVERT")
    print("\n🔄 .xlsb → .xlsx बदल रहे हैं...")
    converted = convert_all_xlsb(files_)
    if not converted:
        print("\n⚠️  कोई .xlsx नहीं बनी।")
        return

    _stage("TRANSFORM")
    transformed, results = transform_all(converted)
    if not transformed:
        print("\n⚠️  कोई फ़ाइल ट्रांसफ़ॉर्म नहीं हुई।")
    else:
        _stage("SHEET")
        push_to_sheets(transformed, results)

    print("\n" + "=" * 60)
    _stage("DONE")
    print("📋 अंतिम रिपोर्ट")
    print("=" * 60)
    for unit, status, note in results:
        print(f"  {unit:<12} {status}  {note}")

    failed = [r for r in results if not r[1].startswith("✅")]
    print("\n" + "=" * 60)
    print("🎉 सब काम पूरा!" if not failed else f"⚠️  काम पूरा, पर {len(failed)} यूनिट में गड़बड़ ऊपर देखें")
    print("=" * 60)


def run():
    global LG
    LG = RunLogger()
    orig = sys.stdout
    sys.stdout = Tee(orig, LG)
    try:
        print(f"🆔 Run ID: {LG.run_id}")
        _run_inner()
        LG.log("COMPLETE", "SUCCESS", stage="DONE", unit="")
    except BaseException as e:
        LG.log(f"FAILED: {e}", "ERROR", stage="DONE", unit="")
        raise
    finally:
        sys.stdout = orig
        LG.close()


# चलाओ
run()
